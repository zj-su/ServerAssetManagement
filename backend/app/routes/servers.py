from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from sqlalchemy.orm import Session
from typing import List, Dict, Any
from app.core.database import get_db
from app.models.server import Server as ServerModel, ServerHistory as ServerHistoryModel
from app.schemas.server import Server, ServerCreate, ServerUpdate
from app.utils.security import get_current_user, require_admin
import logging
import io
from openpyxl import load_workbook

router = APIRouter(prefix="/servers", tags=["servers"])
logger = logging.getLogger(__name__)

# 服务器资产编码前缀与格式：SY-SR-0001
ASSET_CODE_PREFIX = "SY-SR-"
ASSET_CODE_DIGITS = 4


def _next_asset_code(db: Session) -> str:
    """生成下一个服务器资产编码 SY-SR-0001, SY-SR-0002, ..."""
    import re
    servers = db.query(ServerModel).filter(
        ServerModel.asset_code.isnot(None),
        ServerModel.asset_code != "",
    ).all()
    numbers = []
    pattern = re.compile(r"^SY-SR-(\d+)$", re.IGNORECASE)
    for s in servers:
        if s.asset_code and pattern.match(s.asset_code):
            numbers.append(int(pattern.match(s.asset_code).group(1)))
    next_num = (max(numbers) + 1) if numbers else 1
    return f"{ASSET_CODE_PREFIX}{next_num:0{ASSET_CODE_DIGITS}d}"

# 服务器导入模板表头（与 服务器导入模板 一致）
IMPORT_HEADERS = {
    "序列号(SN)": "serial_number",
    "品牌": "brand",
    "型号": "model",
    "BMC IP": "bmc_ip",
    "BMC账号": "bmc_username",
    "BMC密码": "bmc_password",
    "使用人": "user_person",
    "系统IP": "ip_address",
    "系统账号": "system_username",
    "系统密码": "system_password",
    "U位": "u_position",
    "U高度": "u_height",
    "状态": "status",
    "位置": "location",
    "部门": "department",
    "合同号": "contract_number",
    "采购日期": "purchase_date",
    "保修到期": "warranty_expire",
    "备注": "remark",
    # 兼容旧模板
    "主机名": "hostname",
    "IP地址": "ip_address",
    "MAC地址": "mac_address",
    "BMC用户名": "bmc_username",
    "CPU型号": "cpu_model",
    "CPU核心数": "cpu_cores",
    "内存(GB)": "memory_gb",
    "磁盘信息": "disk_info",
}
REQUIRED_FIELDS = ("serial_number", "ip_address")  # 序列号与系统IP至少其一，优先用系统IP去重

# 批量导入：独立路由挂到 /api/v1/import-servers，避免被 /servers/{server_id} 抢匹配导致 405
import_router = APIRouter(tags=["servers"])


@import_router.post("/import-servers/")
def import_servers_excel(
    file: UploadFile = File(..., description="服务器导入模板_(5).xlsx 格式的 Excel"),
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """
    批量导入服务器：按「服务器导入模板_(5).xlsx」表头解析，
    导入后记录出现在「服务器资产 -> 服务器」列表中。
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="请上传 .xlsx 或 .xls 格式的 Excel 文件")
    try:
        content = file.file.read()
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            raise HTTPException(status_code=400, detail="Excel 无有效工作表")
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(status_code=400, detail="Excel 无数据行")
        header_row = [str(c).strip() if c is not None else "" for c in rows[0]]
        col_map = {}
        for idx, h in enumerate(header_row):
            if h in IMPORT_HEADERS:
                col_map[IMPORT_HEADERS[h]] = idx
        if "ip_address" not in col_map and "serial_number" not in col_map:
            raise HTTPException(
                status_code=400,
                detail=f"缺少必填列，请使用模板。需要列：序列号(SN) 或 系统IP；当前表头：{header_row}"
            )
        created = 0
        skipped = []
        errors = []
        for row_idx, row in enumerate(rows[1:], start=2):
            if not row:
                continue
            vals = [row[i] if i < len(row) else None for i in range(max(col_map.values()) + 1)]
            def get(f: str):
                if f not in col_map:
                    return None
                v = vals[col_map[f]]
                if v is None:
                    return None
                if f in ("cpu_cores", "memory_gb", "u_height"):
                    try:
                        return int(float(v))
                    except (TypeError, ValueError):
                        return None
                if f in ("purchase_date", "warranty_expire") and v:
                    try:
                        from datetime import datetime as dt
                        if hasattr(v, "date"):
                            return v.date() if hasattr(v, "date") else v
                        return dt.strptime(str(v).strip()[:10], "%Y-%m-%d").date() if str(v).strip() else None
                    except Exception:
                        return None
                s = str(v).strip() if v else None
                return s if s else None
            ip_address = get("ip_address")
            serial_number = get("serial_number")
            hostname = get("hostname") or serial_number or ip_address or ""
            if not ip_address and not serial_number:
                errors.append({"row": row_idx, "reason": "序列号(SN) 或 系统IP 至少填一项"})
                continue
            if ip_address and db.query(ServerModel).filter(ServerModel.ip_address == ip_address).first():
                skipped.append({"row": row_idx, "ip": ip_address, "reason": "系统IP已存在"})
                continue
            if serial_number and db.query(ServerModel).filter(ServerModel.serial_number == serial_number).first():
                skipped.append({"row": row_idx, "sn": serial_number, "reason": "序列号已存在"})
                continue
            try:
                server_data = {
                    "asset_code": _next_asset_code(db),
                    "serial_number": serial_number,
                    "hostname": hostname,
                    "ip_address": ip_address or "",
                    "mac_address": get("mac_address") or "",
                    "brand": get("brand"),
                    "model": get("model"),
                    "bmc_ip": get("bmc_ip"),
                    "bmc_username": get("bmc_username"),
                    "bmc_password": get("bmc_password"),
                    "user_person": get("user_person"),
                    "system_username": get("system_username"),
                    "system_password": get("system_password"),
                    "u_position": get("u_position"),
                    "u_height": get("u_height"),
                    "status": get("status") or "offline",
                    "location": get("location"),
                    "department": get("department"),
                    "contract_number": get("contract_number"),
                    "purchase_date": get("purchase_date"),
                    "warranty_expire": get("warranty_expire"),
                    "remark": get("remark"),
                    "cpu_model": get("cpu_model"),
                    "cpu_cores": get("cpu_cores"),
                    "memory_gb": get("memory_gb"),
                    "disk_info": get("disk_info"),
                }
                server_data = {k: v for k, v in server_data.items() if v is not None}
                if not server_data.get("status"):
                    server_data["status"] = "offline"
                db_server = ServerModel(**server_data)
                db.add(db_server)
                db.commit()
                db.refresh(db_server)
                created += 1
            except Exception as e:
                db.rollback()
                errors.append({"row": row_idx, "reason": str(e)})
        wb.close()
        return {
            "message": "批量导入完成",
            "created": created,
            "skipped": len(skipped),
            "skipped_details": skipped[:20],
            "errors": errors[:20],
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("服务器批量导入失败: %s", e)
        raise HTTPException(status_code=500, detail=f"导入失败: {str(e)}")


@router.get("/", response_model=List[Server])
def read_servers(
    skip: int = 0, 
    limit: int = 100, 
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    try:
        servers = db.query(ServerModel).offset(skip).limit(limit).all()
        return servers
    except Exception as e:
        logger.error(f"Error fetching servers: {e}")
        raise HTTPException(status_code=500, detail="Internal server error")

@router.post("/", response_model=Server)
def create_server(
    server: ServerCreate, 
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    try:
        data = server.dict(exclude_unset=True)
        # 兼容旧库中 hostname/ip_address/mac_address 为 NOT NULL 的情况
        data.setdefault("hostname", "")
        data.setdefault("ip_address", "")
        data.setdefault("mac_address", "")
        # 资产编码由后端自动生成，规则 SY-SR-0001
        generated_code = _next_asset_code(db)
        data["asset_code"] = generated_code
        db_server = ServerModel(**data)
        db.add(db_server)
        db.commit()
        db.refresh(db_server)
        # 确保响应中一定带上 asset_code（部分环境 refresh 后可能未加载新列）
        if not getattr(db_server, "asset_code", None) or str(db_server.asset_code).strip() == "":
            db_server.asset_code = generated_code
        # 添加创建历史记录
        operator = current_user.username if hasattr(current_user, 'username') else str(current_user)
        _add_server_history(db, db_server.id, "创建服务器", operator, f"创建服务器 {db_server.serial_number or db_server.hostname}")
        return db_server
    except Exception as e:
        logger.error(f"Error creating server: {e}")
        db.rollback()
        raise HTTPException(status_code=500, detail="Internal server error")

@router.get("/{server_id}", response_model=Server)
def read_server(
    server_id: int, 
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    try:
        db_server = db.query(ServerModel).filter(ServerModel.id == server_id).first()
        if db_server is None:
            raise HTTPException(status_code=404, detail="服务器未找到")
        return db_server
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error fetching server {server_id}: {e}")
        raise HTTPException(status_code=500, detail="Internal server error")

# ServerUpdate 允许更新的字段名（与模型列名一致，避免多余/错误键）
_UPDATE_FIELDS = set(ServerUpdate.__fields__.keys())

@router.put("/{server_id}", response_model=Server)
def update_server(
    server_id: int, 
    server: ServerUpdate, 
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    try:
        db_server = db.query(ServerModel).filter(ServerModel.id == server_id).first()
        if db_server is None:
            raise HTTPException(status_code=404, detail="服务器未找到")
        update_data = server.dict(exclude_unset=True)
        for key in list(update_data.keys()):
            if key not in _UPDATE_FIELDS:
                update_data.pop(key, None)
        # 记录变更内容
        changes = []
        for key, value in update_data.items():
            if hasattr(ServerModel, key):
                old_value = getattr(db_server, key)
                if old_value != value:
                    changes.append(f"{key}: {old_value} → {value}")
                setattr(db_server, key, value)
        db.commit()
        db.refresh(db_server)
        # 添加更新历史记录
        if changes:
            operator = current_user.username if hasattr(current_user, 'username') else str(current_user)
            _add_server_history(db, db_server.id, "编辑服务器", operator, "; ".join(changes))
        return db_server
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error updating server {server_id}: {e}")
        db.rollback()
        raise HTTPException(status_code=500, detail="Internal server error")

@router.delete("/{server_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_server(
    server_id: int, 
    db: Session = Depends(get_db),
    current_user = Depends(require_admin)
):
    try:
        db_server = db.query(ServerModel).filter(ServerModel.id == server_id).first()
        if db_server is None:
            raise HTTPException(status_code=404, detail="服务器未找到")
        
        db.delete(db_server)
        db.commit()
        return None
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error deleting server {server_id}: {e}")
        db.rollback()
        raise HTTPException(status_code=500, detail="Internal server error")

# 电源控制相关接口
@router.post("/{server_id}/power/on")
def power_on_server(
    server_id: int, 
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    try:
        db_server = db.query(ServerModel).filter(ServerModel.id == server_id).first()
        if db_server is None:
            raise HTTPException(status_code=404, detail="服务器未找到")
        
        # 这里应该调用实际的BMC/IPMI接口来开机
        # 示例代码，实际需要根据具体硬件实现
        # bmc_control.power_on(db_server.bmc_ip, db_server.bmc_username, db_server.bmc_password)
        
        # 更新服务器状态
        db_server.status = "online"
        db.commit()
        db.refresh(db_server)
        
        return {"message": f"服务器 {db_server.hostname} 已开机", "server": db_server}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error powering on server {server_id}: {e}")
        db.rollback()
        raise HTTPException(status_code=500, detail="Internal server error")

@router.post("/{server_id}/power/off")
def power_off_server(
    server_id: int, 
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    try:
        db_server = db.query(ServerModel).filter(ServerModel.id == server_id).first()
        if db_server is None:
            raise HTTPException(status_code=404, detail="服务器未找到")
        
        # 这里应该调用实际的BMC/IPMI接口来关机
        # 示例代码，实际需要根据具体硬件实现
        # bmc_control.power_off(db_server.bmc_ip, db_server.bmc_username, db_server.bmc_password)
        
        # 更新服务器状态
        db_server.status = "offline"
        db.commit()
        db.refresh(db_server)
        
        return {"message": f"服务器 {db_server.hostname} 已关机", "server": db_server}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error powering off server {server_id}: {e}")
        db.rollback()
        raise HTTPException(status_code=500, detail="Internal server error")

@router.post("/{server_id}/power/reboot")
def reboot_server(
    server_id: int, 
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    try:
        db_server = db.query(ServerModel).filter(ServerModel.id == server_id).first()
        if db_server is None:
            raise HTTPException(status_code=404, detail="服务器未找到")
        
        # 这里应该调用实际的BMC/IPMI接口来重启
        # 示例代码，实际需要根据具体硬件实现
        # bmc_control.reboot(db_server.bmc_ip, db_server.bmc_username, db_server.bmc_password)
        
        # 更新服务器状态
        db_server.status = "maintenance"
        db.commit()
        db.refresh(db_server)
        
        return {"message": f"服务器 {db_server.hostname} 正在重启", "server": db_server}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error rebooting server {server_id}: {e}")
        db.rollback()
        raise HTTPException(status_code=500, detail="Internal server error")


# 获取服务器历史记录
@router.get("/{server_id}/history")
def get_server_history(
    server_id: int,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """获取服务器的操作历史记录"""
    db_server = db.query(ServerModel).filter(ServerModel.id == server_id).first()
    if db_server is None:
        raise HTTPException(status_code=404, detail="服务器未找到")
    
    history = db.query(ServerHistoryModel).filter(
        ServerHistoryModel.server_id == server_id
    ).order_by(ServerHistoryModel.date.desc()).all()
    
    return [
        {
            "id": h.id,
            "server_id": h.server_id,
            "date": h.date.isoformat() if h.date else None,
            "action": h.action,
            "operator": h.operator,
            "remark": h.remark
        }
        for h in history
    ]


# 添加服务器历史记录（内部使用）
def _add_server_history(db: Session, server_id: int, action: str, operator: str, remark: str = None):
    """添加服务器操作历史记录"""
    from datetime import datetime
    import pytz
    now = datetime.now(pytz.timezone('Asia/Shanghai'))
    history = ServerHistoryModel(
        server_id=server_id,
        date=now,
        action=action,
        operator=operator,
        remark=remark
    )
    db.add(history)
    db.commit()